"""다운로드용 통합 문서의 표시 형식과 문자열 보존."""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def format_sheet(ws, header_row=1):
    ws.freeze_panes=f'A{header_row+1}'
    ws.auto_filter.ref=f'A{header_row}:{get_column_letter(ws.max_column)}{max(header_row,ws.max_row)}'
    ws.row_dimensions[header_row].height=32
    for cell in ws[header_row]:
        cell.fill=PatternFill('solid',fgColor='EAF3F4')
        cell.font=Font(bold=True,color='203746')
        cell.alignment=Alignment(vertical='center',wrap_text=True)
    for column in ws.columns:
        letter=column[0].column_letter
        ws.column_dimensions[letter].width=min(38,max(14,max((len(str(c.value or '')) for c in column[:50]),default=12)+2))
        for cell in column:
            if isinstance(cell.value,str) and cell.value.startswith('='):
                cell.data_type='s'
            if cell.row>header_row and isinstance(cell.value,(int,float)):
                cell.number_format='#,##0.##'
