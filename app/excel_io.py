"""엑셀 읽기 유틸. pandas 3는 xlrd 1.2.0과 호환되지 않아 .xls는 xlrd로 직접 읽는다."""

from __future__ import annotations

from datetime import datetime, date
from pathlib import Path

import pandas as pd
import xlrd
from openpyxl import load_workbook


def _cell_to_python(cell: xlrd.sheet.Cell, datemode: int):
    if cell.ctype == xlrd.XL_CELL_EMPTY:
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, datemode)
        except Exception:
            return cell.value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        return cell.value
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return None
    text = str(cell.value).strip()
    return text if text != "" else None


def read_xls_sheet(path: Path | str, sheet_name: str, header_row: int = 2) -> pd.DataFrame:
    book = xlrd.open_workbook(str(path), formatting_info=False)
    sheet = book.sheet_by_name(sheet_name)
    headers = [
        str(sheet.cell_value(header_row, c)).replace("\r\n", "\n")
        for c in range(sheet.ncols)
    ]
    records = []
    for r in range(header_row + 1, sheet.nrows):
        row = [_cell_to_python(sheet.cell(r, c), book.datemode) for c in range(sheet.ncols)]
        records.append(row)
    df = pd.DataFrame(records, columns=headers)
    return df


def read_xlsx_sheet(path: Path | str, sheet_name: str | int | None = 0, header_row: int = 0) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl", header=header_row)


def excel_sheet_names(path: Path | str) -> list[str]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xls":
        book = xlrd.open_workbook(str(path), formatting_info=False)
        return book.sheet_names()
    wb = load_workbook(path, read_only=True, data_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def to_iso_date(value) -> str | None:
    if value is None or (isinstance(value, float) and pd.isna(value)) or pd.isna(value):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date) and not isinstance(value, datetime):
        return value.isoformat()
    ts = pd.to_datetime(value, errors="coerce")
    if pd.isna(ts):
        return None
    return ts.date().isoformat()
