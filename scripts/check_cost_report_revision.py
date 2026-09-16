import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from datetime import date
from io import BytesIO
import importlib.util
from openpyxl import load_workbook
import pandas as pd
from app.db import get_connection
from app.services.cost_report import build_report_xlsx,load_completed

root=Path(__file__).resolve().parents[1]
template=load_workbook(next(p for p in root.glob('*수정필요본.xlsx') if not p.name.startswith('~')))
expected=[c.value for c in template['재작업현황'][1] if c.fill.fgColor.index!='FFFF0000']
conn=get_connection()
start,end=date(2026,9,1),date(2026,9,15)
result=load_workbook(BytesIO(build_report_xlsx(conn,start,end)))
assert [c.value for c in result['재작업현황'][1]]==expected
source=load_completed(conn,start,end)
source=source[source['category']=='완제품']
def clean(value): return None if pd.isna(value) else value
for row,original in zip(list(result['재작업현황'].values)[1:],source.to_dict('records')):
    record=dict(zip(expected,row))
    for target,key in [('SITE','site'),('무/재작업 현황','구분'),('재작업수량','완료수량'),('재작업공수','투입공수')]:
        assert record[target]==clean(original[key]),target
    assert record['MODEL']==str(original['모델명'])[:4]
    assert record['구분']=='생산계획'
assert result['재작업현황'].max_row==len(source)+1
spec=importlib.util.spec_from_file_location('previous_cost_report',root/'data/backups/cost_report_20260915/cost_report.py')
old=importlib.util.module_from_spec(spec);spec.loader.exec_module(old)
previous=load_workbook(BytesIO(old.build_report_xlsx(conn,start,end)))
assert list(result['대여반납'].values)==list(previous['대여반납'].values)
conn.close()
print(f'PASS: {len(source)} finished rows, template headers, six mappings, rental unchanged')
