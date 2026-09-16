import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.stdout.reconfigure(encoding='utf-8')
from app.migration.migrate_master import load_category
from openpyxl import load_workbook
for category, pattern in [('완제품','*완제품*Master.xls'),('대여','*대여*Master.xls')]:
    df, report = load_category(next(Path('.').glob(pattern)), category)
    print(category, {k:report[k] for k in ('after_cut','status_done','status_confirmed')}, 'invalid dates',len(report['date_fail']), 'missing model',df['모델명'].isna().sum(), 'missing inbound',df['입고일'].isna().sum())
for path in Path('.').glob('*.xlsx'):
    wb=load_workbook(path, read_only=True, data_only=True)
    print(path.name, [(s.title,s.max_row,s.max_column) for s in wb])
    for ws in wb:
        print(ws.title, list(ws.values)[:3])
    wb.close()
