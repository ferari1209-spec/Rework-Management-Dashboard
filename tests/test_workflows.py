import io
import sqlite3
import tempfile
import unittest
from datetime import date
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from app.db import get_connection, ensure_schema, seed_settings
from app.data_access import complete_items, update_item_fields, delete_inventory_items
from app.services.import_service import commit_import
from app.services.parsing import split_remark
from app.services.inventory import annotate_inventory
from app.services.cost_report import calc_amount, build_report_xlsx
from app.services.master_export import build_master_xlsx
from app.services.settings_service import set_value
from app.services.upload_parser import detect_and_parse
from app.migration.migrate_master import load_category

class Workflows(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.conn=get_connection(Path(self.tmp.name)/'test.db')
        ensure_schema(self.conn)
        seed_settings(self.conn)
        self.conn.execute("INSERT INTO users(name,login_id,password_hash,role,approved) VALUES('Admin','admin','unused','관리자',1)")
        self.conn.execute("INSERT INTO users(name,login_id,password_hash,role,approved) VALUES('Viewer','viewer','unused','일반',1)")
        self.conn.commit()
    def tearDown(self):
        self.conn.close()
        self.tmp.cleanup()
    def frame(self, **overrides):
        row={'모델명':'A4TC','site':'강남구청','담당자':'담당자','serial':'0001','입고수량':4,'완료수량':0,'투입공수':0.8,'입고일':'2026-01-01','비고_원문':'강남구청,,0001'}
        row.update(overrides)
        return pd.DataFrame([row])
    def save(self, frame=None, digest='test', kind='master'):
        return commit_import(self.conn, self.frame() if frame is None else frame,'완제품',digest,'file.xlsx',date(2026,9,11),'admin',kind)
    def test_parsing_20_cases(self):
        for i in range(20):
            with self.subTest(i=i):
                result=split_remark(f'사업장{i},담당{i},00{i}')
                self.assertEqual((result['site'],result['담당자'],result['serial']), (f'사업장{i}',f'담당{i}',f'00{i}'))
        self.assertEqual(split_remark('사업장,,001')['serial'],'001')
        self.assertFalse(split_remark('사업장,,001')['담당자'])
        self.assertIn('비고 누락',split_remark(float('nan'))['flags'])

    def test_master_sync_update_skip_and_insert(self):
        self.save(self.frame(source_no=1),digest='first')
        old_id=self.conn.execute('SELECT id FROM rework_items').fetchone()[0]
        self.assertEqual(self.save(self.frame(source_no=1),digest='same'),0)
        new=pd.concat([self.frame(source_no=1,완료수량=1,재작업일='2026-09-16',담당팀='생산팀'),
                       self.frame(source_no=2,serial='0002')],ignore_index=True)
        stats=commit_import(self.conn,new,'완제품','changed','master.xls',date.today(),'admin','master',return_summary=True)
        self.assertEqual(stats,{'신규':1,'변경':1,'동일':0})
        rows=self.conn.execute('SELECT * FROM rework_items ORDER BY id').fetchall()
        self.assertEqual(len(rows),2)
        self.assertEqual(rows[0]['id'],old_id)
        self.assertEqual(rows[0]['status'],'완료')
        self.assertEqual(rows[0]['담당팀'],'생산팀')

    def test_master_sync_blank_preserves_and_conflict_rolls_back(self):
        self.save(self.frame(source_no=1,담당팀='원팀'),digest='first')
        self.save(self.frame(source_no=1,투입공수='-',담당팀=None),digest='blank')
        row=self.conn.execute('SELECT * FROM rework_items').fetchone()
        self.assertEqual(row['투입공수'],0.8)
        self.assertEqual(row['담당팀'],'원팀')
        incoming=pd.concat([self.frame(source_no=2,serial='new'),self.frame(source_no=1,serial='wrong')])
        with self.assertRaises(ValueError):self.save(incoming,digest='conflict')
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM rework_items').fetchone()[0],1)
        self.assertFalse(self.conn.execute("SELECT 1 FROM import_batches WHERE digest='conflict'").fetchone())

    def test_master_sync_ambiguous_existing_not_merged(self):
        self.save(self.frame(),digest='first')
        self.save(self.frame(),digest='erp',kind='erp')
        self.conn.execute("UPDATE rework_items SET 입고일='2026-01-01'");self.conn.commit()
        with self.assertRaises(ValueError):self.save(self.frame(완료수량=1),digest='ambiguous')

    def test_upload_editable_fields_and_dash_hours(self):
        from app.services.import_service import editable_import_frame, review_rows
        frame=self.frame(투입공수=' - ',완료수량=1)
        frame['flags']='old';frame['has_issue']=True
        editable=editable_import_frame(frame,'erp',date(2026,9,11))
        self.assertNotIn('flags',editable)
        self.assertTrue(all(str(dtype)=='string' for dtype in editable.dtypes))
        editable.loc[0,'입고일']='2026-09-10'
        editable.loc[0,'비고_원문']='검수 수정'
        reviewed=review_rows(editable,self.conn)
        self.assertNotIn('투입공수',reviewed.iloc[0]['flags'])
        commit_import(self.conn,reviewed,'완제품','dash','file.xls',date(2026,9,11),'admin','erp',reviewed_dates=True)
        row=self.conn.execute('SELECT * FROM rework_items').fetchone()
        self.assertIsNone(row['투입공수'])
        self.assertEqual(row['입고일'],'2026-09-10')
        self.assertEqual(row['비고_원문'],'검수 수정')
        self.assertEqual(row['status'],'완료')
        with self.assertRaises(ValueError):self.save(self.frame(투입공수='abc'),digest='invalid')
    def test_analysis_permissions_and_outliers(self):
        from app.services.analysis_service import analyze
        self.save(pd.concat([self.frame(투입공수=h,완료수량=1,재작업일='2026-09-01',구분='재작업') for h in [1,1,1,1,1,10]],ignore_index=True))
        result=analyze(self.conn,1,'완제품',date(2026,9,1),date(2026,9,30))
        self.assertEqual(result['completed'],6)
        self.assertEqual(len(result['anomalies']),1)
        self.assertEqual(result['anomalies'].iloc[0]['대당공수'],10)
        self.assertEqual(len(result['repeated']),6)
        with self.assertRaises(PermissionError):
            analyze(self.conn,999,'완제품',date(2026,9,1),date(2026,9,30))
        self.conn.execute("UPDATE users SET permissions=? WHERE id=2", ('{"cost":{"view":false}}',))
        self.conn.commit()
        with self.assertRaises(PermissionError):
            analyze(self.conn,2,'완제품',date(2026,9,1),date(2026,9,30))
        self.conn.execute('UPDATE users SET approved=0 WHERE id=1');self.conn.commit()
        with self.assertRaises(PermissionError):
            analyze(self.conn,1,'완제품',date(2026,9,1),date(2026,9,30))
    def test_selected_inventory_delete(self):
        self.save()
        with self.assertRaises(PermissionError):
            delete_inventory_items(self.conn,[1],'완제품','viewer','오입력')
        with self.assertRaises(ValueError):
            delete_inventory_items(self.conn,[1,999],'완제품','admin','오입력')
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM rework_items').fetchone()[0],1)
        self.assertEqual(delete_inventory_items(self.conn,[1],'완제품','admin','오입력'),1)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM rework_items').fetchone()[0],0)
        log = self.conn.execute('SELECT * FROM deleted_item_logs').fetchone()
        self.assertEqual(log['deleted_by'],'admin')
        self.assertIn('A4TC',log['row_json'])
    def test_manager_team_import_and_user_delete(self):
        from app.services.manager_service import fill_teams, delete_users
        self.conn.execute("INSERT INTO managers(name,team,is_active) VALUES('담당자','생산팀',1)")
        self.conn.commit()
        self.save()
        self.assertEqual(self.conn.execute('SELECT 담당팀 FROM rework_items').fetchone()[0],'생산팀')
        frame = fill_teams(self.frame(담당팀='기존팀'), self.conn)
        self.assertEqual(frame.iloc[0]['담당팀'],'기존팀')
        with self.assertRaises(ValueError):
            delete_users(self.conn,[1],'admin')
        with self.assertRaises(PermissionError):
            delete_users(self.conn,[1],'viewer')
        delete_users(self.conn,[2],'admin')
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM users').fetchone()[0],1)
    def test_delete_rejects_completed_inventory(self):
        self.save()
        complete_items(self.conn,[1],date(2026,9,11),None,'admin')
        with self.assertRaises(ValueError):
            delete_inventory_items(self.conn,[1],'완제품','admin','오입력')
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM rework_items').fetchone()[0],1)
    def test_duplicate_and_atomic_validation(self):
        self.save()
        with self.assertRaises(ValueError): self.save()
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM rework_items').fetchone()[0],1)
        with self.assertRaises(ValueError): self.save(pd.concat([self.frame(),self.frame(입고수량=-2)]),digest='bad')
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM rework_items').fetchone()[0],1)
    def test_completion_cost_and_archive(self):
        self.save()
        self.assertEqual(complete_items(self.conn,[1],date(2026,9,11),None,'admin'),1)
        item=dict(self.conn.execute('SELECT * FROM rework_items').fetchone())
        self.assertEqual(item['완료수량'],4)
        self.assertEqual(complete_items(self.conn,[1],date(2026,9,12),None,'admin'),0)
        self.assertFalse(annotate_inventory(pd.DataFrame([item]),self.conn).iloc[0]['is_inventory'])
        wb=load_workbook(io.BytesIO(build_report_xlsx(self.conn,date(2026,9,1),date(2026,9,30))),data_only=True)
        self.assertEqual(wb.sheetnames,['재작업현황','대여반납'])
        headers = [c.value for c in wb.worksheets[0][1]]
        self.assertEqual(wb.worksheets[0].cell(2,headers.index('금액')+1).value,26256)
        record = dict(zip(headers,[c.value for c in wb.worksheets[0][2]]))
        self.assertEqual(record['MODEL'],'A4TC')
        self.assertEqual(record['구분'],'생산계획')
        self.assertEqual(record['SITE'],'강남구청')
        self.assertEqual(record['재작업수량'],4)
        self.assertEqual(record['재작업공수'],0.8)
        self.assertNotIn('사유',headers)
        self.assertEqual(calc_amount(.8,547),26256)
        backup=load_workbook(io.BytesIO(build_master_xlsx(self.conn,'완제품')))
        self.assertEqual(backup.worksheets[0].max_row,4)
    def test_inventory_threshold_and_disposal(self):
        self.save(self.frame(구분='폐기'))
        df=pd.read_sql_query('SELECT * FROM rework_items',self.conn)
        out=annotate_inventory(df,self.conn,date(2026,9,11))
        self.assertTrue(out.iloc[0]['is_inventory'])
        self.assertTrue(out.iloc[0]['is_long_term'])
        set_value(self.conn,'long_term_threshold_days','999','완제품')
        self.assertFalse(annotate_inventory(df,self.conn,date(2026,9,11)).iloc[0]['is_long_term'])
        set_value(self.conn,'excluded_status_values','["폐기"]')
        self.assertFalse(annotate_inventory(df,self.conn).iloc[0]['is_inventory'])
    def test_server_permission(self):
        self.save()
        with self.assertRaises(PermissionError): complete_items(self.conn,[1],date.today(),None,'viewer')
        with self.assertRaises(PermissionError): update_item_fields(self.conn,1,{'모델명':'bad'},'viewer')
        with self.assertRaises(PermissionError): commit_import(self.conn,self.frame(),'완제품','x','x',date.today(),'viewer','erp')
    def test_erp_date_and_master_fields(self):
        self.save(kind='erp')
        self.assertEqual(self.conn.execute('SELECT 입고일 FROM rework_items').fetchone()[0],'2026-09-11')
        self.save(self.frame(재작업일='2026-08-01',변경모델명='NEW'),digest='master')
        row=self.conn.execute('SELECT * FROM rework_items WHERE id=2').fetchone()
        self.assertEqual((row['status'],row['투입공수'],row['변경모델명']),('완료',.8,'NEW'))
    def test_real_master_rows(self):
        for cat,pattern,count,done in [('완제품','*완제품*Master.xls',1511,1508),('대여','*대여*Master.xls',1570,1523)]:
            df,report=load_category(next(Path('.').glob(pattern)),cat)
            self.assertEqual(len(df),count)
            self.assertEqual(report['status_done'],done)
            self.assertFalse(report['date_fail'])
    def test_real_erp(self):
        df,kind=detect_and_parse(Path('대여반납 전산 입고 엑셀 장표.xlsx'),'대여',self.conn)
        self.assertEqual(kind,'erp')
        self.assertEqual(len(df),4)
        self.assertTrue(df['모델명'].notna().all())

    def test_completion_mark_without_hours_or_date(self):
        self.save(self.frame(입고수량=1,완료수량=1,투입공수=None,재작업일=None))
        row=dict(self.conn.execute('SELECT * FROM rework_items').fetchone())
        self.assertEqual(row['status'],'완료')
        self.assertIsNone(row['재작업일'])
        self.assertIsNone(row['투입공수'])
        out=annotate_inventory(pd.DataFrame([row]),self.conn)
        self.assertFalse(out.iloc[0]['is_inventory'])
        self.assertFalse(out.iloc[0]['is_long_term'])
        from app.services.completion import is_completed
        self.assertFalse(is_completed({'투입공수':1,'완료수량':0}))
        self.assertTrue(is_completed({'완료수량':'1'}))

if __name__=='__main__': unittest.main()
