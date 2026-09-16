import unittest
from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from app.services.kpi_details import previous_inventory, detail_excel


class KpiDetailsTests(unittest.TestCase):
    def test_month_end_inventory_and_unknown_completion(self):
        df = pd.DataFrame([
            ['2026-01-01', '2026-09-03', 1, '', '완료'],
            ['2026-01-01', '2026-08-31', 1, '', '완료'],
            ['2026-09-01', None, 0, '', '입고'],
            ['2026-01-01', None, 1, '', '완료'],
            [None, None, 0, '', '입고'],
            ['2026-01-01', None, 0, '폐기', '입고'],
            ['2026-08-31', None, 0, '', '입고'],
        ], columns=['입고일','재작업일','완료수량','구분','status'])
        inventory, long_term, unknown = previous_inventory(df, pd.Period('2026-09'), 90, ['폐기'])
        self.assertEqual(inventory.index.tolist(), [0, 6])
        self.assertEqual(long_term.index.tolist(), [0])
        self.assertEqual(unknown, 2)

    def test_export_keeps_text_and_both_periods(self):
        df = pd.DataFrame({'모델명':['=1+1'], '입고수량':[3]})
        book = load_workbook(BytesIO(detail_excel(df, df.iloc[:0], '집계 안내')))
        self.assertEqual(book.sheetnames, ['현재 상세','전월 상세','집계 기준'])
        self.assertEqual(book['현재 상세']['A2'].data_type, 's')
        self.assertEqual(book['현재 상세']['B2'].value, 3)
